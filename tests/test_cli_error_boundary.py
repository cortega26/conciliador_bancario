from __future__ import annotations

import json
from pathlib import Path

import pytest
from conciliador_bancario.cli import app
from openpyxl import load_workbook
from typer.testing import CliRunner


def _write(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8")


def _write_min_bank_expected(tmp_path: Path) -> tuple[Path, Path]:
    bank = tmp_path / "bank.csv"
    exp = tmp_path / "expected.csv"
    _write(
        bank,
        "\n".join(
            [
                "fecha_operacion,monto,descripcion",
                "05/01/2026,1000,TEST",
                "",
            ]
        ),
    )
    _write(
        exp,
        "\n".join(
            [
                "fecha,monto,descripcion",
                "05/01/2026,1000,TEST",
                "",
            ]
        ),
    )
    return bank, exp


def test_cli_error_config_sin_traceback_por_defecto(tmp_path: Path) -> None:
    cfg = tmp_path / "config.yaml"
    _write(cfg, "cliente: [\n")
    bank, exp = _write_min_bank_expected(tmp_path)

    r = CliRunner().invoke(
        app,
        [
            "validate",
            "--config",
            str(cfg),
            "--bank",
            str(bank),
            "--expected",
            str(exp),
        ],
    )
    assert r.exit_code == 3
    assert "Error (configuracion)" in r.stdout
    assert "Como resolver:" in r.stdout
    assert "Traceback" not in r.stdout


def test_cli_error_config_muestra_traceback_con_debug(tmp_path: Path) -> None:
    cfg = tmp_path / "config.yaml"
    _write(cfg, "cliente: [\n")
    bank, exp = _write_min_bank_expected(tmp_path)

    r = CliRunner().invoke(
        app,
        [
            "validate",
            "--config",
            str(cfg),
            "--bank",
            str(bank),
            "--expected",
            str(exp),
            "--debug",
        ],
    )
    assert r.exit_code == 3
    assert "Traceback" in r.stdout


def test_cli_run_error_ingestion_emite_evento_cli_error(tmp_path: Path) -> None:
    cfg = tmp_path / "config.yaml"
    bank = tmp_path / "bank_bad.csv"
    exp = tmp_path / "expected.csv"
    out = tmp_path / "out"
    out.mkdir()

    _write(cfg, "cliente: 'X'\n")
    _write(
        bank,
        "\n".join(
            [
                "fecha_operacion,monto",
                "05/01/2026,1000",
                "",
            ]
        ),
    )
    _write(exp, "fecha,monto,descripcion\n05/01/2026,1000,TEST\n")

    r = CliRunner().invoke(
        app,
        [
            "run",
            "--config",
            str(cfg),
            "--bank",
            str(bank),
            "--expected",
            str(exp),
            "--out",
            str(out),
            "--dry-run",
        ],
    )
    assert r.exit_code == 4
    assert "CSV banco sin columnas requeridas: descripcion" in r.stdout

    # El fallo de CLI se registra aparte: audit.jsonl es la traza por corrida y
    # toda linea debe poder atribuirse a un run_id, que aqui no existe todavia.
    fallos = out / "audit_fallo.jsonl"
    assert fallos.exists()
    assert any('"tipo":"cli_error"' in line for line in fallos.read_text().splitlines())

    # La traza de la corrida sigue siendo utilizable y esta atribuida.
    audit_path = out / "audit.jsonl"
    assert audit_path.exists()
    for line in audit_path.read_text(encoding="utf-8").splitlines():
        assert "run_id" in json.loads(line)


def test_cli_run_error_auditoria_fallida_no_oculta_error_principal(tmp_path: Path) -> None:
    cfg = tmp_path / "config.yaml"
    bank, exp = _write_min_bank_expected(tmp_path)
    out_file = tmp_path / "out_as_file"
    _write(out_file, "no-dir")
    _write(cfg, "cliente: 'X'\n")

    r = CliRunner().invoke(
        app,
        [
            "run",
            "--config",
            str(cfg),
            "--bank",
            str(bank),
            "--expected",
            str(exp),
            "--out",
            str(out_file),
            "--dry-run",
        ],
    )
    assert r.exit_code == 6
    assert "No se pudo preparar el directorio de salida." in r.stdout


# Cuenta de 10+ digitos: es lo que enmascarar_texto_sensible transforma en
# cuentas enmascaradas, y permite distinguir un reporte con mask de uno sin mask.
_CUENTA_LARGA = "123456789012"

_FLAGS = pytest.mark.parametrize(
    "flags",
    [[], ["--mask"], ["--no-mask"]],
    ids=["sin-flag", "mask", "no-mask"],
)


def _correr_run_con_flags(tmp_path: Path, flags: list[str]) -> Path:
    cfg = tmp_path / "config.yaml"
    bank = tmp_path / "bank.csv"
    exp = tmp_path / "exp.csv"
    out = tmp_path / "out"
    out.mkdir(parents=True)

    _write(cfg, "cliente: 'X'\npermitir_ocr: false\nmoneda_default: 'CLP'\n")
    _write(
        bank,
        "\n".join(
            [
                "fecha_operacion,monto,moneda,descripcion",
                f"05/01/2026,1000,CLP,Transferencia cuenta {_CUENTA_LARGA}",
                "",
            ]
        ),
    )
    _write(exp, "fecha,monto,moneda,descripcion\n05/01/2026,1000,CLP,Pago\n")

    r = CliRunner().invoke(
        app,
        [
            "run",
            "--config",
            str(cfg),
            "--bank",
            str(bank),
            "--expected",
            str(exp),
            "--out",
            str(out),
            *flags,
        ],
    )
    assert r.exit_code == 0, r.stdout
    return out


@_FLAGS
def test_flags_de_masking_terminan_en_cero(tmp_path: Path, flags: list[str]) -> None:
    """--mask, --no-mask y la ausencia de flag son las tres rutas validas."""
    out = _correr_run_con_flags(tmp_path, flags)
    assert (out / "reporte_conciliacion.xlsx").exists()


def test_no_mask_producece_un_reporte_realmente_distinto(tmp_path: Path) -> None:
    """
    --no-mask no basta con aceptarse: tiene que cambiar el reporte.

    Un flag que se parsea y no hace nada tambien exitingia con codigo 0, asi que
    la asercion relevante es que los dos workbooks difieren en el dato sensible.
    """
    out_mask = _correr_run_con_flags(tmp_path / "mask", ["--mask"])
    out_sin = _correr_run_con_flags(tmp_path / "sin", ["--no-mask"])

    def _descripciones(out: Path) -> list[str]:
        ws = load_workbook(out / "reporte_conciliacion.xlsx")["Transacciones"]
        headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        idx = headers.index("descripcion")
        return [str(row[idx]) for row in ws.iter_rows(min_row=2, values_only=True)]

    con_mask = _descripciones(out_mask)
    sin_mask = _descripciones(out_sin)
    assert con_mask != sin_mask
    assert all(_CUENTA_LARGA not in d for d in con_mask)
    assert any(_CUENTA_LARGA in d for d in sin_mask)
